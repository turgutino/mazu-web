"""Tests for document (Word/PDF/Excel/text) attachments on mazu-web's chat
message endpoint. Mirrors test_web_chat_images.py's fixtures -- this file adds
the `files` field to that same POST /api/chat/<id>/message flow. Unlike images
(whose base64 passes straight through unchanged), a file's raw bytes are
extracted to plain text server-side (see mazu/files/document_extract.py)
*before* ever reaching ChatSession, so these tests exercise that extraction
through the real HTTP route, not just the extraction module in isolation
(already covered by tests/test_document_extract.py in the mazu repo).
"""

import base64
import io
import subprocess
import time

import pytest

import mazu_web.chat_session as chat_session_module
from mazu.chat.store import ChatStore
from mazu.llm.types import AgentResponse
from mazu_web.app import create_app


@pytest.fixture()
def project(tmp_path, monkeypatch):
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key-not-real")
    subprocess.run(["git", "init"], cwd=tmp_path, check=True, capture_output=True)
    subprocess.run(["git", "config", "user.email", "test@example.com"], cwd=tmp_path, check=True)
    subprocess.run(["git", "config", "user.name", "Test"], cwd=tmp_path, check=True)
    (tmp_path / ".mazu").mkdir(exist_ok=True)
    return tmp_path


def _end_turn_stream(messages, system, tools, on_delta, model=None):
    on_delta("got the file")
    return AgentResponse(stop_reason="end_turn", content=[{"type": "text", "text": "got the file"}], usage={})


def _wait_for(outbox, event_type, timeout=20):  # see test_web.py's own _wait_for for why 20s
    deadline = time.time() + timeout
    seen = []
    while time.time() < deadline:
        event = outbox.get(timeout=deadline - time.time())
        seen.append(event)
        if event["type"] == event_type:
            return event, seen
    raise AssertionError(f"never saw {event_type}, saw {seen}")


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


def _make_docx_bytes(text: str) -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_text_file_extracted_and_persisted_as_delimited_block(project, monkeypatch):
    monkeypatch.setattr(chat_session_module, "run_turn_stream", _end_turn_stream)
    app = create_app(project, None, None)
    client = app.test_client()
    session_id = client.post("/api/chat/start").get_json()["session_id"]
    session = app.sessions[session_id]

    res = client.post(
        f"/api/chat/{session_id}/message",
        json={
            "text": "what does this say?",
            "files": [{"filename": "notes.txt", "data": _b64(b"the quarterly numbers look good")}],
        },
    )
    assert res.status_code == 200
    _wait_for(session.outbox, "turn_done")
    session.close()
    session.join(timeout=5)

    chat_store = ChatStore(project / ".mazu" / "chat_history.db")
    messages = chat_store.get_messages(session_id)
    chat_store.close()

    content = messages[0]["content"]
    assert content[0]["type"] == "text"
    assert "--- Attached file: notes.txt ---" in content[0]["text"]
    assert "the quarterly numbers look good" in content[0]["text"]
    assert content[1] == {"type": "text", "text": "what does this say?"}


def test_docx_file_extracted_through_the_real_route(project, monkeypatch):
    monkeypatch.setattr(chat_session_module, "run_turn_stream", _end_turn_stream)
    app = create_app(project, None, None)
    client = app.test_client()
    session_id = client.post("/api/chat/start").get_json()["session_id"]
    session = app.sessions[session_id]

    docx_bytes = _make_docx_bytes("Revenue increased by 12 percent this quarter.")
    client.post(
        f"/api/chat/{session_id}/message",
        json={"text": "", "files": [{"filename": "report.docx", "data": _b64(docx_bytes)}]},
    )
    _wait_for(session.outbox, "turn_done")
    session.close()
    session.join(timeout=5)

    chat_store = ChatStore(project / ".mazu" / "chat_history.db")
    messages = chat_store.get_messages(session_id)
    chat_store.close()

    text_block = messages[0]["content"][0]["text"]
    assert "Revenue increased by 12 percent this quarter." in text_block


def _make_pptx_bytes(title: str, body: str) -> bytes:
    import pptx

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text_frame.text = body
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def test_pptx_file_extracted_through_the_real_route(project, monkeypatch):
    monkeypatch.setattr(chat_session_module, "run_turn_stream", _end_turn_stream)
    app = create_app(project, None, None)
    client = app.test_client()
    session_id = client.post("/api/chat/start").get_json()["session_id"]
    session = app.sessions[session_id]

    pptx_bytes = _make_pptx_bytes("Q3 Roadmap", "Ship the new onboarding flow by October.")
    client.post(
        f"/api/chat/{session_id}/message",
        json={"text": "", "files": [{"filename": "roadmap.pptx", "data": _b64(pptx_bytes)}]},
    )
    _wait_for(session.outbox, "turn_done")
    session.close()
    session.join(timeout=5)

    chat_store = ChatStore(project / ".mazu" / "chat_history.db")
    messages = chat_store.get_messages(session_id)
    chat_store.close()

    text_block = messages[0]["content"][0]["text"]
    assert "Q3 Roadmap" in text_block
    assert "Ship the new onboarding flow by October." in text_block


def test_unreadable_file_rejected_with_400(project):
    app = create_app(project, None, None)
    client = app.test_client()
    session_id = client.post("/api/chat/start").get_json()["session_id"]

    garbage = bytes([0xFF, 0xFE, 0x00, 0x01, 0x80, 0x81])
    res = client.post(
        f"/api/chat/{session_id}/message",
        json={"text": "hi", "files": [{"filename": "mystery.bin", "data": _b64(garbage)}]},
    )
    assert res.status_code == 400
    assert "mystery.bin" in res.get_json()["error"]

    app.sessions[session_id].close()


def test_corrupted_docx_rejected_with_400(project):
    app = create_app(project, None, None)
    client = app.test_client()
    session_id = client.post("/api/chat/start").get_json()["session_id"]

    res = client.post(
        f"/api/chat/{session_id}/message",
        json={"text": "hi", "files": [{"filename": "broken.docx", "data": _b64(b"not a real docx")}]},
    )
    assert res.status_code == 400
    assert "broken.docx" in res.get_json()["error"]

    app.sessions[session_id].close()


def test_missing_filename_rejected_with_400(project):
    app = create_app(project, None, None)
    client = app.test_client()
    session_id = client.post("/api/chat/start").get_json()["session_id"]

    res = client.post(
        f"/api/chat/{session_id}/message",
        json={"text": "hi", "files": [{"data": _b64(b"hello")}]},
    )
    assert res.status_code == 400

    app.sessions[session_id].close()


def test_oversized_file_rejected_with_400(project):
    app = create_app(project, None, None)
    client = app.test_client()
    session_id = client.post("/api/chat/start").get_json()["session_id"]

    huge_data = "A" * 20_000_001
    res = client.post(
        f"/api/chat/{session_id}/message",
        json={"text": "hi", "files": [{"filename": "big.txt", "data": huge_data}]},
    )
    assert res.status_code == 400
    assert "too large" in res.get_json()["error"]

    app.sessions[session_id].close()


def test_no_files_field_still_works_as_before(project, monkeypatch):
    # Backward compatibility: existing clients that never send `files` at all
    # must be completely unaffected.
    monkeypatch.setattr(chat_session_module, "run_turn_stream", _end_turn_stream)
    app = create_app(project, None, None)
    client = app.test_client()
    session_id = client.post("/api/chat/start").get_json()["session_id"]
    session = app.sessions[session_id]

    client.post(f"/api/chat/{session_id}/message", json={"text": "hi"})
    _wait_for(session.outbox, "turn_done")
    session.close()
    session.join(timeout=5)

    chat_store = ChatStore(project / ".mazu" / "chat_history.db")
    messages = chat_store.get_messages(session_id)
    chat_store.close()
    assert messages[0] == {"role": "user", "content": "hi"}
